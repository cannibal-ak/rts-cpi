"""Shared CSV/XLSX row parsing helpers used by the ingestion service.

These helpers were previously inlined in ``scripts/ingest_daily.py``; they
are kept here so the working logic survives the Phase 6 removal of that
script. The helpers are intentionally lenient: bad rows return safe
defaults (``0``, ``None``) and the caller decides whether the row is
valid based on critical fields.
"""
from __future__ import annotations

import csv
import logging
from datetime import date, datetime, time
from typing import Iterator

# App-level loggers don't propagate in this container (see main.py:27);
# route through "uvicorn.error" so the auto-detect notice reaches stderr.
log = logging.getLogger("uvicorn.error")

# Canonical 13-column header for VELOCITY CSV files. When a velocity
# upload arrives without a header row (some airline exports omit it),
# read_data_file injects these names so DictReader yields rows keyed
# the same way as a headed file. Order matches the on-disk column
# layout — do not reorder without confirming exporter output.
_VELOCITY_CANONICAL_HEADERS: list[str] = [
    "DepDate",
    "DepTime",
    "DepCode",
    "CityPair",
    "Eqp",
    "LegsegType",
    "LegSegOrder",
    "Days_Left",
    "Compartment",
    "Current_Booking",
    "Capacity",
    "Actual_Seat_Factor",
    "Forecasted_Seat_Factor",
]


def parse_date(value: str | None) -> date | None:
    """Parse a date string in any of the formats produced by the source files.

    Returns ``None`` if the value is missing, blank, or unparseable.
    """
    if not value:
        return None
    s = str(value).strip()
    if not s or s == "0":
        return None
    for fmt in ("%d%b%y", "%d%B%y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except (ValueError, TypeError):
            continue
    try:
        return datetime.strptime(s.upper(), "%d%b%y").date()
    except (ValueError, TypeError):
        return None


def parse_velocity_date(value: str | None) -> date | None:
    """Parse a date string from a velocity CSV (M/D/YYYY or ISO)."""
    if not value:
        return None
    s = str(value).strip()
    if not s or s == "0":
        return None
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except (ValueError, TypeError):
            continue
    return None


def parse_time(value: str | None) -> time | None:
    """Parse a time string. Returns ``None`` if missing or unparseable."""
    if not value:
        return None
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M%p"):
        try:
            return datetime.strptime(s, fmt).time()
        except (ValueError, TypeError):
            continue
    return None


def safe_float(value: object) -> float:
    if not value:
        return 0.0
    try:
        return float(str(value).replace(",", "").strip() or 0)
    except (ValueError, TypeError):
        return 0.0


def safe_float_nullable(value: object) -> float | None:
    """Like ``safe_float``, but returns ``None`` for missing/blank/unparseable input.

    Use this for nullable numeric DB columns where 0.0 would be a valid
    semantic value (e.g. a missing return-leg fare should be NULL, not
    0.0 — which would distort averages and filtering). Also treats the
    literal strings "NA", "N/A", and "NULL" (case-insensitive) as None.
    """
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.upper() in {"NA", "N/A", "NULL"}:
        return None
    try:
        return float(s.replace(",", ""))
    except (ValueError, TypeError):
        return None


def safe_int(value: object) -> int:
    if not value:
        return 0
    try:
        return int(str(value).strip() or 0)
    except (ValueError, TypeError):
        return 0


def safe_int_nullable(value: object) -> int | None:
    """Like ``safe_int``, but returns ``None`` for missing/blank/unparseable input.

    Use this for nullable integer DB columns where 0 would be a valid
    semantic value (e.g. ``ref_stops = 0`` is "nonstop"; we don't want
    a missing value to silently look like a nonstop flight).
    """
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.upper() == "NULL":
        return None
    try:
        return int(s)
    except (ValueError, TypeError):
        return None


def read_data_file(
    file_path: str, domain: str | None = None
) -> Iterator[dict]:
    """Yield rows from a CSV or XLSX file as ``dict``s keyed by header.

    Auto-detects format by extension. The XLSX path uses openpyxl in
    read-only mode so memory stays bounded for large workbooks.

    When ``domain == "VELOCITY"`` and the file is CSV, peeks at line 1
    and injects the canonical 13-column header if the file is detected
    as headerless. Other domains and the XLSX path are unchanged.
    """
    if file_path.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                return
            headers = [
                str(h).strip() if h is not None else f"col_{i}"
                for i, h in enumerate(header_row)
            ]
            for row_values in rows_iter:
                row_dict: dict[str, str] = {}
                for i, val in enumerate(row_values):
                    if i < len(headers):
                        row_dict[headers[i]] = (
                            str(val).strip() if val is not None else ""
                        )
                yield row_dict
        finally:
            wb.close()
        return

    inject_velocity_header = False
    if domain == "VELOCITY":
        with open(file_path, "r", encoding="utf-8-sig") as fh:
            first_line = fh.readline()
        if first_line:
            first_row = next(csv.reader([first_line]), [])
            first_field = first_row[0].strip() if first_row else ""
            n_fields = len(first_row)
            if first_field != "DepDate":
                if n_fields != 13:
                    raise ValueError(
                        f"Headerless velocity file has {n_fields} "
                        "columns, expected 13"
                    )
                inject_velocity_header = True
                log.info(
                    "Auto-detected headerless velocity file; injecting "
                    "canonical 13-column header (file=%s)",
                    file_path,
                )

    with open(file_path, "r", encoding="utf-8-sig") as fh:
        if inject_velocity_header:
            reader = csv.DictReader(
                fh, fieldnames=_VELOCITY_CANONICAL_HEADERS
            )
        else:
            reader = csv.DictReader(fh)
        for row in reader:
            yield row

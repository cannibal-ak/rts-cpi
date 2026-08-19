#!/usr/bin/env python3
"""Build the DreamAir demo dataset from PW's source files.

DreamAir's demo data is PW's feed re-badged. This script produces the DA_* files;
it does NOT ingest them (see ingest_dreamair_demo.py).

By default every PW file on disk is processed. The archive runs from 2025-08-16
to 2026-08-19, though the 18 airline files before 2025-09-03 use an older layout
the parser rejects ("Unrecognized AIRLINE schema"), and velocity does not start
until 2025-12-12 -- so Sep-Nov 2025 are fare-only days. Use --start/--end to
narrow the window; the original build ran --start 2026-01-01 --end 2026-07-01.

WHY RE-INGEST RATHER THAN COPY THE ROWS
---------------------------------------
PW's rows in the database are damaged. Every Jan-Jun departure and arrival time
is truncated to four characters ("20:4", not "20:45") and 1.28M flight numbers
are cut at ten ("PW0423/PW0"), because they were ingested before migration 037
widened those columns. The source .xlsx files are intact -- RefDepTime reads
"05:00" -- so re-ingesting recovers what an INSERT ... SELECT from PW would
carry over broken. Two WinAir-parity charts depend on those fields:

  * Latest Prices' departure-time and duration sliders parse "HH:MM"; "20:4"
    matches neither that nor the "HHMM" fallback.
  * Deployed Capacity unnests ref_flt_num on "/", so "PW0423/PW0" contributes a
    phantom leg "PW0".

WHAT IS REWRITTEN
-----------------
  RefAL      PW     -> DA
  RefFltNum  PW0423 -> DA0423, per leg, so "PW0423/PW0718" -> "DA0423/DA0718"

Nothing else. Competitor columns are left alone: TC, KQ, Fli, Aur, YS, Coa, UI
and CQ are genuinely other airlines, and PW never appears as its own
competitor (verified: 0 rows).

Velocity CSVs are copied byte-for-byte and only renamed. They are headerless
and carry no airline column at all -- the tenant comes from the filename -- so
DA_VL_080126.csv is deliberately identical to PW_VL_080126.csv. That is why the
dedup gate in ingestion/service.py is tenant-scoped; without that, this file
would be rejected as a duplicate of PW's committed job.

DUPLICATE SOURCES
-----------------
A few basenames exist in more than one staging directory (re-uploads). The copy
whose SHA-256 matches a COMMITTED PW job is authoritative, so the DreamAir data
matches what PW actually serves. If none matches, the largest is used and the
choice is logged.

Usage (inside the api container, which has openpyxl and the DB URL):
    python3 scripts/relabel_pw_files_as_da.py --out /tmp/da_files
    python3 scripts/relabel_pw_files_as_da.py --out /tmp/da_files --limit 3   # smoke test
"""
from __future__ import annotations

import argparse
import hashlib
import os
import re
import shutil
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl
import sqlalchemy as sa

# Paths as the api container sees them: the repo's apps/api is bind-mounted at
# /app, and the repo root at /sample_data. Non-existent roots are skipped, so
# this is also safe to run on the host with the paths adjusted.
SRC_ROOTS = [
    Path("/app/data/staging"),
    Path("/sample_data/dev-sftp-mount"),
]

# DDMMYY, matching filename_patterns.yaml
AIRLINE_RE = re.compile(r"^PW_(\d{2})(\d{2})(\d{2})\.xlsx$", re.IGNORECASE)
VELOCITY_RE = re.compile(r"^PW_VL_(\d{2})(\d{2})(\d{2})\.(csv|xlsx)$", re.IGNORECASE)

SRC_CODE = "PW"
DST_CODE = "DA"


def _file_date(m: re.Match) -> date | None:
    dd, mm, yy = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        return date(2000 + yy, mm, dd)
    except ValueError:
        return None


def _sha256(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _committed_pw_hashes(engine) -> set[str]:
    with engine.connect() as conn:
        rows = conn.execute(sa.text(
            "SELECT file_hash FROM ingestion_jobs "
            "WHERE tenant_code = :tc AND status = 'COMMITTED'"
        ), {"tc": SRC_CODE})
        return {r[0] for r in rows}


def discover(start: date | None, end: date | None) -> dict[str, list[Path]]:
    """basename -> every copy of it found on disk. Bounds are [start, end)."""
    found: dict[str, list[Path]] = defaultdict(list)
    for root in SRC_ROOTS:
        if not root.exists():
            continue
        for p in root.rglob("*"):
            if not p.is_file():
                continue
            m = AIRLINE_RE.match(p.name) or VELOCITY_RE.match(p.name)
            if not m:
                continue
            d = _file_date(m)
            if d is None:
                continue
            if (start is not None and d < start) or (end is not None and d >= end):
                continue
            found[p.name].append(p)
    return found


def pick(copies: list[Path], committed: set[str]) -> tuple[Path, str]:
    """Choose the authoritative copy. Returns (path, reason)."""
    if len(copies) == 1:
        return copies[0], "only copy"
    for p in copies:
        if _sha256(p) in committed:
            return p, "hash matches a COMMITTED PW job"
    best = max(copies, key=lambda p: p.stat().st_size)
    return best, f"NO committed match among {len(copies)}; chose largest"


def relabel_airline(src: Path, dst: Path) -> tuple[int, int]:
    """Rewrite RefAL and RefFltNum in place into dst. Returns (rows, flt_rewrites)."""
    wb = openpyxl.load_workbook(src)
    ws = wb.active

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        i_al = header.index("RefAL")
    except ValueError:
        raise SystemExit(f"{src.name}: no RefAL column; headers={header[:12]}")
    i_flt = header.index("RefFltNum") if "RefFltNum" in header else None

    rows = 0
    flt_rewrites = 0
    for row in ws.iter_rows(min_row=2):
        rows += 1
        cell = row[i_al]
        if isinstance(cell.value, str) and cell.value.strip().upper() == SRC_CODE:
            cell.value = DST_CODE
        if i_flt is not None:
            fc = row[i_flt]
            if isinstance(fc.value, str) and fc.value:
                # per leg, so multi-leg itineraries are fully rebadged
                legs = fc.value.split("/")
                new = "/".join(
                    DST_CODE + leg[len(SRC_CODE):]
                    if leg.upper().startswith(SRC_CODE) else leg
                    for leg in legs
                )
                if new != fc.value:
                    fc.value = new
                    flt_rewrites += 1

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    wb.close()
    return rows, flt_rewrites


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True, help="output directory for DA_* files")
    ap.add_argument("--limit", type=int, default=0, help="process at most N of each kind")
    ap.add_argument("--force", action="store_true", help="rewrite outputs that already exist")
    ap.add_argument("--start", type=date.fromisoformat, default=None,
                    help="earliest file date to process, YYYY-MM-DD (default: no bound)")
    ap.add_argument("--end", type=date.fromisoformat, default=None,
                    help="exclusive upper bound on file date, YYYY-MM-DD (default: no bound)")
    args = ap.parse_args()

    if args.start and args.end and args.start >= args.end:
        raise SystemExit(f"--start {args.start} is not before --end {args.end}")

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    url = os.environ.get("CPI_DATABASE_URL") or os.environ.get("DATABASE_URL")
    if not url:
        raise SystemExit("CPI_DATABASE_URL not set")
    engine = sa.create_engine(url)
    committed = _committed_pw_hashes(engine)
    print(f"committed PW file hashes: {len(committed)}")

    found = discover(args.start, args.end)
    airline = sorted(n for n in found if AIRLINE_RE.match(n))
    velocity = sorted(n for n in found if VELOCITY_RE.match(n))
    window = f"{args.start or 'beginning'}..{args.end or 'end'}"
    print(f"discovered in {window}: "
          f"{len(airline)} airline, {len(velocity)} velocity")

    if args.limit:
        airline = airline[: args.limit]
        velocity = velocity[: args.limit]

    n_air = n_vel = 0
    total_rows = total_flt = 0

    for name in airline:
        src, reason = pick(found[name], committed)
        if "NO committed match" in reason:
            print(f"  ! {name}: {reason}")
        dst = out / ("DA_" + name[len("PW_"):])
        if dst.exists() and not args.force:
            n_air += 1
            continue
        rows, flt = relabel_airline(src, dst)
        total_rows += rows
        total_flt += flt
        n_air += 1
        if n_air % 20 == 0:
            print(f"  airline {n_air}/{len(airline)} ...", flush=True)

    for name in velocity:
        src, reason = pick(found[name], committed)
        if "NO committed match" in reason:
            print(f"  ! {name}: {reason}")
        dst = out / ("DA_VL_" + name[len("PW_VL_"):])
        if dst.exists() and not args.force:
            n_vel += 1
            continue
        shutil.copyfile(src, dst)
        n_vel += 1

    print(f"\nwrote {n_air} airline files ({total_rows} data rows, "
          f"{total_flt} flight numbers rebadged) and {n_vel} velocity files to {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
